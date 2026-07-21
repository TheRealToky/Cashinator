"""Structural check of a generated export against the shop's legacy format.

Run after `flutter test`, which writes `build/test_exports/sample_export.xlsx`:

    python tool/verify_export.py [path/to/reference.xlsx]

Compares the generated workbook against the reference ledger export
(`pastry_sales_2026-06-19.xlsx` by default) on the things the downstream
automation script depends on: sheet names, header text and order, cell types,
and the TOTAL row formulas. Exits non-zero on any mismatch.
"""

from __future__ import annotations

import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install openpyxl")

GENERATED = Path("build/test_exports/sample_export.xlsx")
MANIFEST = Path("build/test_exports/sample_export.expected.txt")

DEFAULT_REFERENCE = Path(
    r"D:\PC DISAINE\toky\Perso-D\red\sales\img_extraction\2026-06"
    r"\pastry_sales_2026-06-19.xlsx"
)

SALES_LINES = "Sales Lines"
DAY_SUMMARY = "Day Summary"

failures: list[str] = []
checks = 0


def check(condition: bool, message: str) -> None:
    global checks
    checks += 1
    if condition:
        print(f"  PASS  {message}")
    else:
        print(f"  FAIL  {message}")
        failures.append(message)


def header_of(worksheet) -> list:
    return [cell for cell in next(worksheet.iter_rows(max_row=1, values_only=True))]


def main() -> int:
    if not GENERATED.exists():
        return fail_fast(f"{GENERATED} not found — run `flutter test` first.")

    reference_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_REFERENCE
    if not reference_path.exists():
        return fail_fast(f"reference workbook not found: {reference_path}")

    generated = openpyxl.load_workbook(GENERATED)
    reference = openpyxl.load_workbook(reference_path)

    print(f"generated : {GENERATED}")
    print(f"reference : {reference_path}\n")

    print("Sheets")
    check(
        generated.sheetnames == [SALES_LINES, DAY_SUMMARY],
        f"sheet names are {[SALES_LINES, DAY_SUMMARY]} "
        f"(got {generated.sheetnames})",
    )

    print("\nSales Lines header")
    gen_header = header_of(generated[SALES_LINES])
    ref_header = header_of(reference[SALES_LINES])
    check(
        gen_header == ref_header,
        "header matches the reference exactly",
    )
    if gen_header != ref_header:
        print(f"        reference: {ref_header}")
        print(f"        generated: {gen_header}")

    print("\nSales Lines cell types")
    sheet = generated[SALES_LINES]
    rows = [r for r in sheet.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    check(bool(rows), "at least one data row was written")

    if rows:
        line_ids = [r[0] for r in rows]
        check(
            len(line_ids) == len(set(line_ids)),
            "Line Id values are unique",
        )
        check(
            line_ids == list(range(1, len(line_ids) + 1)),
            "Line Id is a 1-based sequence",
        )
        check(
            all(isinstance(r[2], str) for r in rows),
            "Order Label is text, as in the reference",
        )
        check(
            all(isinstance(r[6], str) and len(r[6]) == 5 and r[6][2] == ":" for r in rows),
            "Time is HH:MM text (not coerced to a time serial)",
        )
        check(
            all(isinstance(r[7], str) and len(r[7]) == 10 for r in rows),
            "Date is YYYY-MM-DD text",
        )
        check(all(r[8] == 1 for r in rows), "Source Image is hardcoded to 1")
        check(
            all(r[11] == "medium" for r in rows),
            "Confidence is hardcoded to 'medium'",
        )
        check(
            all(r[10] in ("normal", "voided") for r in rows),
            "Status is 'normal' or 'voided'",
        )
        check(
            all(isinstance(r[9], (int, float)) for r in rows),
            "Total Price is numeric",
        )

        voided = [r for r in rows if r[10] == "voided"]
        check(bool(voided), "the sample contains a voided row to check")
        check(
            all(r[12] for r in voided),
            "voided rows carry the void reason in Notes",
        )
        normal = [r for r in rows if r[10] == "normal"]
        check(
            all(not r[12] for r in normal),
            "normal rows leave Notes blank",
        )

    print("\nDay Summary")
    summary = generated[DAY_SUMMARY]
    gen_summary_header = header_of(summary)
    ref_summary_header = header_of(reference[DAY_SUMMARY])

    check(
        gen_summary_header[:4] == ref_summary_header[:4],
        "leading columns match the reference",
    )
    check(
        gen_summary_header[-1] == ref_summary_header[-1] == "Low Confidence Lines",
        "trailing column is 'Low Confidence Lines'",
    )
    check(
        gen_summary_header[4:7] == ["MoMo", "Card", "Cash"],
        "the three legacy payment columns keep their reference order",
    )

    summary_rows = [
        r for r in summary.iter_rows(min_row=2, values_only=True) if r[0] is not None
    ]
    day_rows = [r for r in summary_rows if r[0] != "TOTAL"]
    total_rows = [r for r in summary_rows if r[0] == "TOTAL"]

    check(bool(day_rows), "at least one day row")
    check(
        [r[0] for r in day_rows] == list(range(1, len(day_rows) + 1)),
        "Source Image is a 1-based day index",
    )
    check(all(r[-1] == 0 for r in day_rows), "Low Confidence Lines is 0")

    payment_columns = gen_summary_header[4:-1]
    for row in day_rows:
        paid = sum(row[4 : 4 + len(payment_columns)])
        check(
            paid == row[3],
            f"day {row[0]}: payment columns sum to Total Sales ({paid} == {row[3]})",
        )

    check(len(total_rows) == 1, "exactly one TOTAL row")
    if total_rows:
        total = total_rows[0]
        formulas = [c for c in total[1:] if isinstance(c, str) and c.startswith("=SUM(")]
        check(
            len(formulas) == len(gen_summary_header) - 1,
            "every TOTAL cell is a SUM() formula",
        )
        expected_last = len(day_rows) + 1
        check(
            all(f.endswith(f"{expected_last})") for f in formulas),
            f"SUM() ranges cover rows 2..{expected_last}",
        )

    print(f"\n{checks - len(failures)}/{checks} checks passed")
    if failures:
        print("\nFAILED:")
        for failure in failures:
            print(f"  - {failure}")
        return 1

    print("Generated export is structurally compatible with the legacy format.")
    return 0


def fail_fast(message: str) -> int:
    print(f"ERROR: {message}")
    return 2


if __name__ == "__main__":
    raise SystemExit(main())
